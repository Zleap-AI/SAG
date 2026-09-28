"""Run with the pinned fnos/develop API environment, NOT the current API.

Only the model/embedding response boundaries are deterministic. The old
parser, relational schema, event saver, and LanceDB writer run unchanged.
"""

import argparse
import asyncio
import hashlib
import importlib.metadata
import json
import sqlite3
import tarfile
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from zleap.sag import DataEngine, EngineConfig
from zleap.sag.config import EmbeddingConfig, LLMConfig
from zleap.sag.core.ai.embedding import EmbeddingClient
from zleap.sag.modules.extract.extractor import EventExtractor
from zleap.sag.modules.extract.processor import EventProcessor
from zleap.sag.modules.load.processor import DocumentProcessor

from sag_api.core.config import Settings
from sag_api.db.base import Base
from sag_api.db.models import Agent, Document, Job, Message, Setting, Source, Thread, User
from sag_api.enums import DocumentStatus, JobStatus, JobType, MessageRole
from sag_api.parsing import prepare_document

BASELINE = "8ad465446d89067337b7406bff09fe44438f2059"
PREFIX = "/fixture-native"
TENANTS = ("1000", "1001-bob-81b637d8")


async def fixed_embedding(self, text):
    return [1.0] + [0.0] * 7


async def fixed_client(self):
    return SimpleNamespace()


async def fixed_batch_embedding(self, texts):
    return [await fixed_embedding(self, text) for text in texts]


async def fixed_response(self, items, metadata, source_type):
    return {
        "data": {
            "items": [
                {
                    "title": "Retained knowledge",
                    "summary": "The original private document must survive upgrade.",
                    "content": "The original private document must survive upgrade.",
                    "references": [1],
                    "entities": [{"type": "concept", "name": "Migration", "description": "Private knowledge"}],
                }
            ]
        }
    }


async def build(root):
    assert importlib.metadata.version("zleap-sag") == "0.7.1"
    assert "knowledge_state" not in Document.__table__.columns
    DocumentProcessor.generate_embedding = fixed_embedding
    EmbeddingClient.generate = fixed_embedding
    EmbeddingClient.batch_generate = fixed_batch_embedding
    EventExtractor._get_llm_client = fixed_client
    EventProcessor.process = fixed_response
    for tenant in TENANTS:
        workspace = root / "users" / tenant
        uploads = workspace / "uploads" / f"source-{tenant}"
        uploads.mkdir(parents=True)
        original = uploads / "book.md"
        original.write_text(f"# Tenant {tenant}\n\n" + "The original private document must survive upgrade. " * 10)
        spreadsheet = uploads / "records.xlsx"
        workbook = Workbook()
        workbook.active.append(["Record", "Owner", "Amount"])
        workbook.active.append(["R001", tenant, 100])
        workbook.active.append([])
        workbook.active.append(["Record", "Owner", "Amount"])
        workbook.active.append(["R002", tenant, 200])
        workbook.save(spreadsheet)
        canonical = await prepare_document(str(spreadsheet), Settings(_env_file=None, document_parser="markitdown"))
        engine = DataEngine(
            EngineConfig(
                data_dir=str(workspace / "engine"),
                language="en",
                llm=LLMConfig(api_key="fixture-only", model="fixture"),
                embedding=EmbeddingConfig(model="fixture", dimensions=8),
            ),
            source_config_id=f"config-{tenant}",
        )
        async with engine:
            ingested = await engine.ingest(original)
            extracted = await engine.extract()
            assert ingested.chunk_count > 0 and extracted.event_count > 0
            sheet_ingested = await engine.ingest(canonical.path)
            sheet_extracted = await engine.extract()
            assert sheet_ingested.chunk_count > 0 and sheet_extracted.event_count > 0

        meta = workspace / "meta" / "sag.db"
        meta.parent.mkdir()
        database = create_async_engine(f"sqlite+aiosqlite:///{meta}")
        async with database.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        sessions = async_sessionmaker(database, expire_on_commit=False)
        async with sessions() as session:
            source = Source(
                id=f"source-{tenant}",
                name=f"Knowledge {tenant}",
                sag_source_config_id=f"config-{tenant}",
                document_count=3,
                chunk_count=ingested.chunk_count + sheet_ingested.chunk_count,
                event_count=extracted.event_count + sheet_extracted.event_count,
            )
            user = User(id=f"user-{tenant}", email=f"{tenant}@fixture.test", password_hash="fixture-only")
            agent = Agent(id=f"agent-{tenant}", name="Retained agent", persona={"greeting": "hello"})
            session.add_all([source, user, agent])
            await session.flush()
            document = Document(
                id=f"doc-{tenant}",
                source_id=source.id,
                filename="book.md",
                storage_path=str(original).replace(str(root), PREFIX),
                status=DocumentStatus.READY,
                sag_source_id=ingested.source_id,
                chunk_count=ingested.chunk_count,
                event_count=extracted.event_count,
            )
            missing = Document(
                id=f"missing-{tenant}",
                source_id=source.id,
                filename="missing.md",
                storage_path=str(uploads / "missing.md").replace(str(root), PREFIX),
                status=DocumentStatus.READY,
            )
            thread = Thread(id=f"thread-{tenant}", agent_id=agent.id, title="Retained conversation")
            sheet = Document(
                id=f"sheet-{tenant}",
                source_id=source.id,
                filename="records.xlsx",
                storage_path=str(spreadsheet).replace(str(root), PREFIX),
                status=DocumentStatus.READY,
                sag_source_id=sheet_ingested.source_id,
                chunk_count=sheet_ingested.chunk_count,
                event_count=sheet_extracted.event_count,
            )
            session.add_all([document, missing, sheet, thread])
            await session.flush()
            session.add_all(
                [
                    Job(
                        id=f"job-{tenant}",
                        source_id=source.id,
                        document_id=document.id,
                        type=JobType.PROCESS_DOCUMENT,
                        status=JobStatus.QUEUED,
                        payload={
                            "process_checkpoint": {"source_id": ingested.source_id, "chunk_ids": ingested.chunk_ids}
                        },
                    ),
                    Message(
                        id=f"message-{tenant}",
                        thread_id=thread.id,
                        role=MessageRole.ASSISTANT,
                        content="Retained answer [1]",
                        citations=[
                            {"n": 1, "chunk_id": ingested.chunk_ids[0], "source_id": source.id},
                            {"n": 2, "kind": "external", "url": "https://example.com"},
                        ],
                    ),
                    Setting(scope="global", key="fixture_config", value={"model": "retained", "tenant": tenant}),
                ]
            )
            await session.commit()
        await database.dispose()
        # Resolve WAL into the files that the stopped Native service backs up.
        for sqlite_file in workspace.rglob("*.db"):
            with sqlite3.connect(sqlite_file) as connection:
                connection.execute("PRAGMA wal_checkpoint(TRUNCATE)")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    asyncio.run(build(args.root))
    files = {
        str(p.relative_to(args.root)): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(args.root.rglob("*"))
        if p.is_file()
    }
    args.output.mkdir(parents=True, exist_ok=True)
    with tarfile.open(args.output / "native-071.tar.gz", "w:gz") as archive:
        archive.add(args.root / "users", arcname="users")
    (args.output / "provenance.json").write_text(
        json.dumps(
            {
                "baseline": BASELINE,
                "engine_version": "0.7.1",
                "path_prefix": PREFIX,
                "tenants": TENANTS,
                "files": files,
                "model_boundary": "Deterministic embedding and extraction responses; real old storage writers",
            },
            indent=2,
        )
        + "\n"
    )


if __name__ == "__main__":
    main()
